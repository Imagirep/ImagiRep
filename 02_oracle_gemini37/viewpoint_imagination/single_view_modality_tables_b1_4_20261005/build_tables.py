"""Rebuild the two single-view modality tables (net change / rescued + best modality per type) for Batch1-4."""
import sys, os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
sys.path.insert(0, str(Path(__file__).resolve().parent)); from sv622 import table622, stats, CATS, C
OUT = Path(__file__).resolve().parent / 'single_view_622_modality_tables.xlsx'
F = 'Arial'; GREEN = PatternFill('solid', fgColor='E6F4EA'); RED = PatternFill('solid', fgColor='FCE8E6'); HEAD = PatternFill('solid', fgColor='EFEFEF')
NAMES = ['Start-only', 'Final (base)', 'Topdown', 'Text', 'Video', 'Depth', 'Normals', 'Panorama']
SRC = ('Source: raw per-question results. B1 selected171 runs (validation / endpoint_modality / endpoint_geometry 2026-09-16); '
       'B2 fo_batch2_eight_conditions_20260927; B3/B4 start_view+start_final from 2026-09-21/24 runs, six modalities from '
       'fo_batch3/4_six_modality_conditions_20261004. Gemini 3.7 flash, temp 0, HIGH thinking. Category object_selection_by_direction excluded (as in the B1+B2 308 table). '
       'Note: B1 Panorama uses the 2026-09-16 pre-repair panorama.')


def cell(ws, r, c, v, bold=False, fill=None, fmt=None, color=None, align=None):
    x = ws.cell(r, c, v); x.font = Font(name=F, bold=bold, color=color)
    if fill: x.fill = fill
    if fmt: x.number_format = fmt
    if align: x.alignment = Alignment(horizontal=align)
    return x


def sheet1(ws, o, title):
    cell(ws, 1, 1, title)
    hdr = ['Category', 'N'] + NAMES
    for j, h in enumerate(hdr, 1): cell(ws, 2, j, h, bold=True, fill=HEAD, align='center')
    for i, (k, lab) in enumerate(CATS):
        v = o[k]; r = 3 + i
        cell(ws, r, 1, lab); cell(ws, r, 2, v['n'], color='0000FF')
        for j, c in enumerate(C):
            n = v['net'][c]; cell(ws, r, 3 + j, n, color='0000FF', fill=None if c == 'start_final' else GREEN if n > 0 else RED if n < 0 else None)
    cell(ws, 10, 1, 'Net change in correct answers vs Start + Final (modality correct - final correct).', color='666666')
    cell(ws, 12, 1, 'By Question')
    for j, h in enumerate(hdr, 1):
        if h != 'Start-only': cell(ws, 13, j, h, bold=True, fill=HEAD, align='center')
    for i, (k, lab) in enumerate(CATS):
        v = o[k]; r = 14 + i; best = max(v['rescue'][c] for c in C[2:])
        cell(ws, r, 1, lab); cell(ws, r, 2, v['n'], color='0000FF'); cell(ws, r, 4, 0, color='0000FF')
        for j, c in enumerate(C[2:]):
            n = v['rescue'][c]; cell(ws, r, 5 + j, n, bold=n == best, color='0000FF', fill=GREEN if n > 0 else None)
    cell(ws, 21, 1, 'Rescued = questions wrong with Start + Final but correct with Start + Final + modality (positive-only). Bold = best per row.', color='666666')
    cell(ws, 23, 1, SRC, color='666666')
    ws.column_dimensions['A'].width = 30
    for col in 'BCDEFGHIJ': ws.column_dimensions[col].width = 13


def sheet2(ws, o, title, s1):
    cell(ws, 1, 1, title, bold=True)
    hdr = ['Category', 'Best Additional Modality', 'Endpoint Correct', 'Rescued Questions', 'Final Correct', 'Total', 'Accuracy Gain (pp)']
    for j, h in enumerate(hdr, 1): cell(ws, 2, j, h, bold=True, fill=HEAD, align='center')
    for i, (k, lab) in enumerate(CATS):
        v = o[k]; r = 3 + i; src = 14 + i
        best = max(v['rescue'][c] for c in C[2:]); names = [NAMES[2 + j] for j, c in enumerate(C[2:]) if v['rescue'][c] == best]
        cell(ws, r, 1, lab); cell(ws, r, 2, names[0] if len(names) == 1 else ' or '.join(names) + ' (tied)')
        c3 = cell(ws, r, 3, v['final'], color='0000FF'); c3.comment = Comment('Start + Final correct count from raw results (input).', 'build_tables.py')
        cell(ws, r, 4, f"=MAX('{s1}'!E{src}:J{src})")
        cell(ws, r, 5, f'=C{r}+D{r}'); cell(ws, r, 6, f"='{s1}'!B{src}")
        cell(ws, r, 7, f'=D{r}/F{r}*100', fmt='0.00" pp"')
    t = 10
    cell(ws, t, 1, 'Total', bold=True, fill=HEAD); cell(ws, t, 2, None, fill=HEAD)
    for j, col in zip(range(3, 7), 'CDEF'): cell(ws, t, j, f'=SUM({col}3:{col}8)', bold=True, fill=HEAD)
    cell(ws, t, 7, f'=D{t}/F{t}*100', bold=True, fill=HEAD, fmt='0.00" pp"')
    cell(ws, 12, 1, 'Upper bound accuracy:'); cell(ws, 12, 2, f'=E{t}/F{t}', fmt='0.00%')
    cell(ws, 13, 1, 'Start + Endpoint baseline:'); cell(ws, 13, 2, f'=C{t}/F{t}', fmt='0.00%')
    cell(ws, 14, 1, 'Improvement:', color='008000'); cell(ws, 14, 2, f'=G{t}', fmt='0.00" pp"', color='008000')
    cell(ws, 16, 1, 'Best modality = most rescued questions per category, selected on the same data (retrospective upper bound).', color='666666')
    ws.column_dimensions['A'].width = 30; ws.column_dimensions['B'].width = 26
    for col in 'CDEFG': ws.column_dimensions[col].width = 17


wb = Workbook(); wb.remove(wb.active)
qs, look, cat, batch = table622()
for tag, lab, sel in [('622', 'Batch1-4 (622 set)', qs), ('622 B1+B2', 'Batch1 + Batch2 within 622', [q for q in qs if batch[q] in (1, 2)])]:
    o = stats(sel, look, cat); n = len(sel)
    s1 = f'{tag} Net & Rescued'; s2 = f'{tag} Best Modality'
    sheet1(wb.create_sheet(s1), o, f'{lab}, {n} questions')
    sheet2(wb.create_sheet(s2), o, f'{lab}, Select best modality for each question type ({n} questions):', s1)
wb.save(OUT); print(OUT)
